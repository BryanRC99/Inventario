import unicodedata
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO

from django.conf import settings
from django.db.models.functions import Lower
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from proveedores.models import Proveedor

from .models import Activo, Categoria, Ubicacion


class ArchivoImportacionInvalido(Exception):
    """El archivo no se puede importar (formato, tamaño o plantilla incorrecta).
    El mensaje es apto para mostrarlo tal cual al usuario."""


COLUMNAS = [
    'Nombre',
    'Categoria',
    'Ubicacion',
    'Marca',
    'Modelo',
    'Numero de serie',
    'Estado',
    'Proveedor',
    'Fecha adquisicion (AAAA-MM-DD)',
    'Valor adquisicion',
    'Fecha fin garantia (AAAA-MM-DD)',
]

# Un activo "dado de baja" exige motivo y cierre de custodias: se hace con la opción
# "Dar de baja", no al importar.
ESTADOS_IMPORTABLES = {
    valor: etiqueta
    for valor, etiqueta in Activo.Estado.choices
    if valor != Activo.Estado.DADO_DE_BAJA
}

LARGO_MAX = {'nombre': 150, 'marca': 100, 'modelo': 100, 'numero_serie': 100}
VALOR_MAX = Decimal('10000000000')  # límite de DecimalField(max_digits=12, decimal_places=2)
FILAS_VACIAS_TOLERADAS = 2000  # filas vacías (con formato) que se toleran leer después de las llenas


def _limpiar(valor):
    if valor is None:
        return ''
    # Excel guarda los números como decimales: 123456 llega como 123456.0.
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _normalizar(texto):
    texto = _limpiar(texto).lower()
    texto = unicodedata.normalize('NFD', texto)
    texto = ''.join(c for c in texto if unicodedata.category(c) != 'Mn')
    return ' '.join(texto.split())  # colapsa espacios múltiples


def _clave_estado(texto):
    """'En mantenimiento', 'en_mantenimiento' y 'EN MANTENIMIENTO' son el mismo estado."""
    return _normalizar(texto).replace('_', ' ')


_ESTADO_POR_CLAVE = {}
for _valor, _etiqueta in ESTADOS_IMPORTABLES.items():
    _ESTADO_POR_CLAVE[_clave_estado(_valor)] = _valor
    _ESTADO_POR_CLAVE[_clave_estado(_etiqueta)] = _valor

_ENCABEZADOS_ESPERADOS = [_normalizar(c) for c in COLUMNAS]


def generar_plantilla_excel():
    wb = Workbook()

    # --- Hoja 1: plantilla para llenar ---
    hoja = wb.active
    hoja.title = 'Activos'
    hoja.append(COLUMNAS)

    for celda in hoja[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='111827')

    # Fila de ejemplo, se puede borrar antes de llenar
    hoja.append(
        [
            'Laptop Dell Latitude 5440', 'Laptop', 'Oficina Sistemas', 'Dell', 'Latitude 5440',
            'SN123456', 'activo', 'TechSupply S.A.', '2026-01-15', '950.00', '2028-01-15',
        ]
    )

    for columna in hoja.columns:
        max_len = max(len(str(c.value)) for c in columna if c.value)
        hoja.column_dimensions[columna[0].column_letter].width = max(14, max_len + 2)

    # Validación desplegable para la columna Estado (columna G)
    dv_estado = DataValidation(
        type='list', formula1=f'"{",".join(ESTADOS_IMPORTABLES.keys())}"', allow_blank=False
    )
    hoja.add_data_validation(dv_estado)
    dv_estado.add(f'G2:G{settings.IMPORTACION_MAX_FILAS + 1}')

    # --- Hoja 2: instrucciones ---
    instrucciones = wb.create_sheet('Instrucciones')
    instrucciones.append(['Instrucciones para llenar la plantilla'])
    instrucciones['A1'].font = Font(bold=True, size=13)
    filas_instrucciones = [
        '',
        '1. No modifiques ni reordenes los encabezados de la hoja "Activos".',
        '2. Borra la fila de ejemplo antes de subir el archivo.',
        '3. "Nombre", "Categoria" y "Ubicacion" son obligatorios. Los demás son opcionales.',
        '4. "Categoria" y "Ubicacion" deben escribirse EXACTAMENTE como aparecen en las listas de abajo.',
        '5. El código interno del activo se genera automáticamente, no lo incluyas.',
        '6. Las fechas van en formato AAAA-MM-DD, ej. 2026-03-20.',
        '7. "Valor adquisicion" solo números, sin símbolo de moneda ni separador de miles, ej. 950.00.',
        f'8. Máximo {settings.IMPORTACION_MAX_FILAS} activos por archivo.',
        '9. No se repite un número de serie: si ya existe en el sistema o está dos veces en el archivo, esa fila se rechaza.',
        '10. Para dar de baja un activo usa la opción "Dar de baja" del sistema; no se puede importar como dado de baja.',
        '11. Al importar elegirás la empresa a la que pertenecen todos los activos del archivo.',
        '',
        'Valores válidos para "Estado" (si lo dejas vacío se usa "activo"):',
    ]
    for linea in filas_instrucciones:
        instrucciones.append([linea])

    for estado_valor, estado_label in ESTADOS_IMPORTABLES.items():
        instrucciones.append([f'  - {estado_valor}  ({estado_label})'])

    instrucciones.append([''])
    instrucciones.append(['Categorías existentes:'])
    for cat in Categoria.objects.order_by('nombre'):
        instrucciones.append([f'  - {cat.nombre}'])

    instrucciones.append([''])
    instrucciones.append(['Ubicaciones existentes:'])
    for ub in Ubicacion.objects.order_by('nombre'):
        instrucciones.append([f'  - {ub.nombre}'])

    instrucciones.append([''])
    instrucciones.append(['Proveedores existentes (opcional):'])
    for prov in Proveedor.objects.order_by('nombre'):
        instrucciones.append([f'  - {prov.nombre}'])

    instrucciones.column_dimensions['A'].width = 60

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _parsear_fecha(valor):
    if valor is None or _limpiar(valor) == '':
        return None, None
    if isinstance(valor, datetime):
        return valor.date(), None
    if isinstance(valor, date):
        return valor, None
    texto = _limpiar(valor)
    try:
        return datetime.strptime(texto, '%Y-%m-%d').date(), None
    except ValueError:
        return None, f'Fecha inválida "{texto}", usa el formato AAAA-MM-DD'


def _parsear_valor(valor):
    """Devuelve (Decimal | None, error | None)."""
    if valor is None or _limpiar(valor) == '':
        return None, None
    if isinstance(valor, (int, float, Decimal)) and not isinstance(valor, bool):
        texto = str(valor)
    else:
        texto = _limpiar(valor).replace(' ', '')
        if ',' in texto and '.' in texto:
            return None, f'Valor de adquisición "{valor}": escríbelo sin separador de miles, ej. 1250.50'
        texto = texto.replace(',', '.')
    try:
        numero = Decimal(texto)
    except InvalidOperation:
        return None, f'Valor de adquisición "{valor}" no es un número válido'
    if not numero.is_finite():
        return None, f'Valor de adquisición "{valor}" no es un número válido'
    if numero < 0:
        return None, 'El valor de adquisición no puede ser negativo'
    numero = numero.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    if numero >= VALOR_MAX:
        return None, 'El valor de adquisición es demasiado grande'
    return numero, None


def _leer_filas(archivo):
    """Lee el Excel (solo lectura, en flujo) y devuelve [(número_de_fila, valores)].
    Solo conserva las filas con algo escrito y falla si el archivo no es la plantilla
    o tiene más filas que el máximo permitido."""
    maximo = settings.IMPORTACION_MAX_FILAS
    try:
        wb = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        raise ArchivoImportacionInvalido(
            'No se pudo leer el archivo. Verifica que sea un Excel (.xlsx) descargado de la plantilla.'
        )

    try:
        hoja = wb['Activos'] if 'Activos' in wb.sheetnames else wb.active
        filas = hoja.iter_rows(values_only=True)

        encabezado = next(filas, None)
        if encabezado is None:
            raise ArchivoImportacionInvalido('El archivo está vacío.')
        encontrados = [_normalizar(c) for c in (list(encabezado) + [None] * len(COLUMNAS))[: len(COLUMNAS)]]
        if encontrados != _ENCABEZADOS_ESPERADOS:
            raise ArchivoImportacionInvalido(
                'El archivo no tiene los encabezados de la plantilla. Descarga la plantilla de nuevo '
                'y no modifiques ni reordenes la primera fila de la hoja "Activos".'
            )

        resultado = []
        leidas = 0
        for idx, fila in enumerate(filas, start=2):
            leidas += 1
            if leidas > maximo + FILAS_VACIAS_TOLERADAS:
                raise ArchivoImportacionInvalido(
                    f'El archivo tiene demasiadas filas. El máximo es {maximo} activos por archivo.'
                )
            if fila is None or all(c is None or _limpiar(c) == '' for c in fila):
                continue  # fila vacía, se ignora silenciosamente
            resultado.append((idx, fila))
            if len(resultado) > maximo:
                raise ArchivoImportacionInvalido(
                    f'El archivo tiene más de {maximo} activos. Divídelo en varios archivos.'
                )
        if not resultado:
            raise ArchivoImportacionInvalido('El archivo no tiene ninguna fila con datos.')
        return resultado
    finally:
        wb.close()


def validar_archivo(archivo):
    """
    Lee el Excel y devuelve una lista de dicts, uno por fila, cada uno con
    sus datos originales, si es válida, y la lista de errores si no.
    No toca la base de datos todavía (solo la consulta).

    Lanza ArchivoImportacionInvalido si el archivo en conjunto no sirve.
    """
    crudas = _leer_filas(archivo)

    categorias = {_normalizar(c.nombre): c for c in Categoria.objects.all()}
    ubicaciones = {_normalizar(u.nombre): u for u in Ubicacion.objects.all()}
    proveedores = {_normalizar(p.nombre): p for p in Proveedor.objects.all()}

    # Números de serie que ya existen en el sistema (una sola consulta para todo el archivo).
    series_archivo = {
        _limpiar(f[5]).lower() for _, f in crudas if len(f) > 5 and _limpiar(f[5])
    }
    series_existentes = set()
    if series_archivo:
        series_existentes = set(
            Activo.objects.exclude(numero_serie='')
            .annotate(serie_minuscula=Lower('numero_serie'))
            .filter(serie_minuscula__in=series_archivo)
            .values_list('serie_minuscula', flat=True)
        )

    filas_resultado = []
    series_vistas = {}

    for idx, fila in crudas:
        (
            nombre, categoria_nombre, ubicacion_nombre, marca, modelo, numero_serie,
            estado, proveedor_nombre, fecha_adq, valor_adq, fecha_garantia,
        ) = (list(fila) + [None] * len(COLUMNAS))[: len(COLUMNAS)]

        errores = []
        datos = {
            'nombre': _limpiar(nombre),
            'marca': _limpiar(marca),
            'modelo': _limpiar(modelo),
            'numero_serie': _limpiar(numero_serie),
        }

        if not datos['nombre']:
            errores.append('El nombre es obligatorio')
        for campo, etiqueta in (
            ('nombre', 'El nombre'), ('marca', 'La marca'),
            ('modelo', 'El modelo'), ('numero_serie', 'El número de serie'),
        ):
            if len(datos[campo]) > LARGO_MAX[campo]:
                errores.append(f'{etiqueta} no puede superar los {LARGO_MAX[campo]} caracteres')

        cat_key = _normalizar(categoria_nombre)
        if not cat_key:
            errores.append('La categoría es obligatoria')
        elif cat_key not in categorias:
            errores.append(f'La categoría "{_limpiar(categoria_nombre)}" no existe')
        else:
            datos['categoria'] = categorias[cat_key]

        ubi_key = _normalizar(ubicacion_nombre)
        if not ubi_key:
            errores.append('La ubicación es obligatoria')
        elif ubi_key not in ubicaciones:
            errores.append(f'La ubicación "{_limpiar(ubicacion_nombre)}" no existe')
        else:
            datos['ubicacion'] = ubicaciones[ubi_key]

        estado_txt = _limpiar(estado)
        if not estado_txt:
            datos['estado'] = Activo.Estado.ACTIVO
        elif _clave_estado(estado_txt) in _ESTADO_POR_CLAVE:
            datos['estado'] = _ESTADO_POR_CLAVE[_clave_estado(estado_txt)]
        elif _clave_estado(estado_txt) == _clave_estado(Activo.Estado.DADO_DE_BAJA):
            datos['estado'] = Activo.Estado.ACTIVO
            errores.append('No se puede importar un activo como "dado de baja"; usa la opción "Dar de baja"')
        else:
            datos['estado'] = Activo.Estado.ACTIVO
            errores.append(f'Estado "{estado_txt}" no es válido')

        prov_key = _normalizar(proveedor_nombre)
        if prov_key:
            if prov_key not in proveedores:
                errores.append(f'El proveedor "{_limpiar(proveedor_nombre)}" no existe')
            else:
                datos['proveedor'] = proveedores[prov_key]

        fecha_adq_parseada, error_fecha1 = _parsear_fecha(fecha_adq)
        if error_fecha1:
            errores.append(f'Fecha de adquisición: {error_fecha1}')
        datos['fecha_adquisicion'] = fecha_adq_parseada

        fecha_gar_parseada, error_fecha2 = _parsear_fecha(fecha_garantia)
        if error_fecha2:
            errores.append(f'Fecha de garantía: {error_fecha2}')
        datos['fecha_fin_garantia'] = fecha_gar_parseada

        if fecha_adq_parseada and fecha_gar_parseada and fecha_gar_parseada < fecha_adq_parseada:
            errores.append('La garantía no puede terminar antes de la fecha de adquisición')

        valor, error_valor = _parsear_valor(valor_adq)
        if error_valor:
            errores.append(error_valor)
        elif valor is not None:
            datos['valor_adquisicion'] = valor

        serie = datos['numero_serie'].lower()
        if serie:
            if serie in series_existentes:
                errores.append(f'El número de serie "{datos["numero_serie"]}" ya está registrado en el sistema')
            elif serie in series_vistas:
                errores.append(
                    f'El número de serie "{datos["numero_serie"]}" está repetido '
                    f'(también en la fila {series_vistas[serie]})'
                )
            else:
                series_vistas[serie] = idx

        filas_resultado.append(
            {
                'fila': idx,
                'valido': len(errores) == 0,
                'errores': errores,
                'datos_mostrar': {
                    'nombre': datos['nombre'],
                    'categoria': _limpiar(categoria_nombre),
                    'ubicacion': _limpiar(ubicacion_nombre),
                    'marca': datos['marca'],
                    'modelo': datos['modelo'],
                },
                'datos': datos,
            }
        )

    return filas_resultado
