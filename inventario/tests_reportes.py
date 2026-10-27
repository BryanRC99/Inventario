from io import BytesIO

from django.template.loader import render_to_string
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from empresas.models import Empresa
from inventario.models import Activo, Categoria, Ubicacion
from usuarios.models import Usuario

URL = '/api/inventario/activos/'


class ReportesActivosTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.e1 = Empresa.objects.create(nombre='Antoisafood')
        cls.e2 = Empresa.objects.create(nombre='Jasafood')
        cat = Categoria.objects.create(nombre='Laptops', prefijo='LAP')
        ub = Ubicacion.objects.create(nombre='Bomboli', tipo='sede')
        cls.admin = Usuario.objects.create_user('adm', password='x', rol='admin')
        Activo.objects.create(codigo_interno='LAP-001', categoria=cat, nombre='Laptop A', ubicacion=ub,
                              empresa=cls.e1, numero_serie='SN-111', marca='Dell')
        Activo.objects.create(codigo_interno='LAP-002', categoria=cat, nombre='Laptop B', ubicacion=ub,
                              empresa=cls.e2, numero_serie='SN-222', marca='HP')
        Activo.objects.create(codigo_interno='LAP-003', categoria=cat, nombre='Laptop sin empresa', ubicacion=ub,
                              empresa=None)

    def setUp(self):
        self.client.force_authenticate(self.admin)

    def hoja(self, params=None):
        r = self.client.get(f'{URL}reporte_excel/', params or {})
        self.assertEqual(r.status_code, 200)
        return load_workbook(BytesIO(r.content))['Activos']

    def test_excel_incluye_la_columna_empresa(self):
        hoja = self.hoja()
        encabezado = [c.value for c in hoja[1]]
        self.assertEqual(encabezado[:4], ['Código', 'Nombre', 'Empresa', 'Categoría'])
        filas = {f[0].value: f[2].value for f in hoja.iter_rows(min_row=2)}
        self.assertEqual(filas, {'LAP-001': 'Antoisafood', 'LAP-002': 'Jasafood', 'LAP-003': '—'})

    def test_excel_filtrado_por_empresa_solo_trae_esa_empresa(self):
        hoja = self.hoja({'empresa': str(self.e2.id)})
        self.assertEqual([(f[0].value, f[2].value) for f in hoja.iter_rows(min_row=2)], [('LAP-002', 'Jasafood')])

    def test_pdf_se_genera_y_la_plantilla_no_usa_letra_monoespaciada(self):
        r = self.client.get(f'{URL}reporte_pdf/')
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b'%PDF'))
        html = render_to_string('inventario/reporte_activos_pdf.html', {
            'filas': [{'codigo': 'LAP-001', 'nombre': 'x', 'empresa': 'Antoisafood', 'categoria': 'c', 'marca': 'm',
                       'modelo': 'm', 'numero_serie': 'SN-1', 'ubicacion': 'u', 'estado': 'Activo', 'custodio': 'c'}],
            'logo_uri': None, 'total': 1, 'filtros': [], 'fecha_generacion': 'hoy',
        })
        self.assertNotIn('monospace', html)
        self.assertIn('<th>Empresa</th>', html)
        self.assertIn('<td>Antoisafood</td>', html)
